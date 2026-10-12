import os
import subprocess
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

def create_html_cv():
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Roddneil B. Gemina - Curriculum Vitae</title>
<style>
  @page {
    size: A4 portrait;
    margin: 9mm 11mm 9mm 11mm;
  }
  
  * {
    box-sizing: border-box;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
  }
  
  body {
    font-family: "Garamond", "EB Garamond", "Times New Roman", serif;
    font-size: 9.8pt;
    line-height: 1.23;
    color: #000000;
    margin: 0;
    padding: 0;
    background: #ffffff;
  }

  /* Header */
  .header {
    text-align: center;
    margin-bottom: 5pt;
  }
  
  .header h1 {
    font-size: 20pt;
    font-weight: bold;
    margin: 0 0 2pt 0;
    letter-spacing: 0.2pt;
  }
  
  .header .contact-line {
    font-size: 10pt;
    margin: 0 0 2pt 0;
  }
  
  .header .links-line {
    font-size: 10pt;
    margin: 0;
  }
  
  .header a, .entry a {
    color: #1155cc;
    text-decoration: underline;
  }
  
  /* Sections */
  .section-title {
    font-size: 11pt;
    font-weight: bold;
    text-transform: uppercase;
    margin-top: 8pt;
    margin-bottom: 3.5pt;
    border-bottom: 1px solid #000000;
    padding-bottom: 1pt;
    letter-spacing: 0.4pt;
  }
  
  .entry {
    margin-bottom: 5pt;
    page-break-inside: avoid;
    break-inside: avoid;
  }
  
  .entry-header {
    font-size: 10.2pt;
    margin-bottom: 1pt;
  }
  
  .entry-header b {
    font-weight: bold;
  }
  
  .entry-date {
    font-size: 9.8pt;
    font-style: italic;
    margin-bottom: 1.5pt;
    color: #111111;
  }
  
  .row-between {
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    font-size: 10.2pt;
  }
  
  .row-between .title-left {
    font-weight: bold;
  }
  
  .row-between .date-right {
    font-style: normal;
    text-align: right;
    white-space: nowrap;
    font-size: 9.8pt;
  }

  /* Bullet Points */
  ul.bullets {
    list-style-type: none;
    margin: 0;
    padding-left: 13pt;
  }
  
  ul.bullets li {
    position: relative;
    font-size: 9.6pt;
    line-height: 1.22;
    margin-bottom: 2pt;
    text-align: justify;
    text-justify: inter-word;
  }
  
  ul.bullets li::before {
    content: "■";
    position: absolute;
    left: -12pt;
    font-size: 6.8pt;
    top: 1.2pt;
    color: #000000;
  }

  /* Paragraph style description under competition/research */
  .entry-desc {
    font-size: 9.6pt;
    line-height: 1.23;
    margin: 1.5pt 0 3.5pt 0;
    text-align: justify;
    text-justify: inter-word;
  }

  /* Technical skills */
  .skills-group {
    font-size: 9.5pt;
    line-height: 1.25;
    margin-bottom: 2.8pt;
    text-align: justify;
  }
  
  .skills-group b {
    font-weight: bold;
  }

  /* Certifications table */
  .cert-row {
    display: flex;
    justify-content: space-between;
    font-size: 9.8pt;
    line-height: 1.32;
    margin-bottom: 1.2pt;
    page-break-inside: avoid;
    break-inside: avoid;
  }
  
  .cert-name {
    font-weight: bold;
  }
  
  .cert-date {
    white-space: nowrap;
    text-align: right;
  }
</style>
</head>
<body>

  <!-- Header -->
  <div class="header">
    <h1>Roddneil B. Gemina</h1>
    <div class="contact-line">
      roddneilgemina@gmail.com &nbsp;|&nbsp; +63 956-241-3577 &nbsp;|&nbsp; Cebu City, PH
    </div>
    <div class="links-line">
      github.com/RoddneilGemina &nbsp;|&nbsp; linkedin.com/in/roddneil-gemina-20b423409
    </div>
  </div>

  <!-- Professional Experience -->
  <div class="section-title">PROFESSIONAL EXPERIENCE</div>

  <div class="entry">
    <div class="entry-header">
      <b>Opswerks – Site Reliability Engineering (SRE) Intern</b> | Cebu City, PH
    </div>
    <div class="entry-date">January 2026 – Present</div>
    <ul class="bullets">
      <li>Administered enterprise Linux server environments, configuring AWS cloud infrastructure, containerized microservices via Docker, and production-grade Kubernetes pod orchestration.</li>
      <li>Constructed automated deployment pipelines and integrated cluster telemetry, distributed logging, and health monitoring tools to sustain high platform reliability and zero-downtime operations.</li>
      <li>Automated repetitive operational runbooks and system provisioning using Python and Bash scripting, significantly reducing manual incident mitigation time and operational toil.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>DXForm Technologies – Quality Assurance Analyst & Software Engineer</b> | Cebu City, PH
    </div>
    <div class="entry-date">January 2025 – August 2025</div>
    <ul class="bullets">
      <li><b>Gigi POS (Next.js, React Native, PostgreSQL, Stripe):</b> Directed core QA testing strategies, authoring structured test matrices and regression suites across web and mobile POS terminals; contributed directly to UI/UX design, application logic, and product experience while developing code-level bug fixes and ensuring transactional integrity for Stripe payment processing.</li>
      <li><b>Aicquire (React, Node.js, MongoDB, AI-Assisted Tooling):</b> Spearheaded frontend UI/UX architecture and responsive components in React, engineered backend application logic and MongoDB data pipelines in Node.js, and refined product experience to streamline AI-assisted tooling interactions.</li>
      <li>Logged reproducible defect lifecycles with root-cause analysis in Jira, collaborated cross-functionally with dev leads to optimize REST API response payloads, and enforced cross-platform UX fidelity.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>Nexchef – Mobile Architect & Full-Stack Developer (Contract / Freelance)</b> | Cebu City, PH
    </div>
    <div class="entry-date">November 2024 – March 2025</div>
    <ul class="bullets">
      <li>Developed a full-stack culinary mobile application in React Native and TypeScript featuring modular recipe creation, live step-by-step guidance, and synchronized multi-stage culinary timer engines.</li>
      <li>Integrated Supabase Realtime and WebSockets for instant state synchronization across multiple client devices, backed by PostgreSQL relational persistence.</li>
      <li>Designed responsive multi-screen layouts with micro-interactions, local offline persistence, and low-latency client-side caching.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>Refertoire – Systems Architect & Full-Stack Developer (Contract / Freelance)</b> | Cebu City, PH
    </div>
    <div class="entry-date">January 2026 – Present</div>
    <ul class="bullets">
      <li>Architected a 100% offline-capable music repertoire management platform, allowing performing ensembles to organize, search, and view complete digital sheet music libraries without internet dependency.</li>
      <li>Engineered local offline persistence utilizing IndexedDB and modern Service Workers, coupled with bidirectional client synchronization protocols upon network reconnection.</li>
      <li>Implemented responsive high-resolution document viewing and instant multi-parameter filtering across hundreds of musical scores.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>DishCover – Full-Stack Developer & Database Architect (Contract / Freelance)</b> | Cebu City, PH
    </div>
    <div class="entry-date">May 2024 – October 2024</div>
    <ul class="bullets">
      <li>Developed a comprehensive pantry and household inventory management application featuring automated expiration tracking, replenishment alerts, and storage categorization.</li>
      <li>Engineered normalized PostgreSQL relational schemas optimized for high-volume item lookup, barcode data association, and historical consumption analytics.</li>
      <li>Built a responsive, minimalist UI with Tailwind CSS and Next.js server components, optimizing page loads and delivering sub-second interaction speed.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>CropConnect – Full-Stack Agritech Developer (Contract / Freelance)</b> | Cebu City, PH
    </div>
    <div class="entry-date">February 2024 – August 2024</div>
    <ul class="bullets">
      <li>Co-developed an agricultural direct-to-consumer e-commerce platform connecting regional farmers directly with commercial buyers to eliminate intermediary price gouging.</li>
      <li>Engineered catalog browsing, inventory batching, and order fulfillment tracking pipelines with PostgreSQL backend models and responsive UI flows.</li>
    </ul>
  </div>

  <!-- Relevant Projects -->
  <div class="section-title" style="page-break-before: always;">RELEVANT PROJECTS</div>

  <div class="entry">
    <div class="entry-header">
      <b>ResuTex – Lead Architect & Full-Stack Developer</b> | Next.js 16, Python, ReportLab, Gemini Flash, Supabase, REST APIs
    </div>
    <ul class="bullets">
      <li>Architected an autonomous ATS resume studio combining Python ReportLab and LaTeX engines to dynamically compile and format tailored resumes and cover letters.</li>
      <li>Engineered dynamic ATS match scoring and qualification extraction algorithms against job postings, integrating LLM-assisted tailoring pipelines with Google Gemini Flash.</li>
      <li>Integrated Supabase REST APIs and local failover ledgers for real-time application tracking, status telemetry, and automated cold outreach drafting.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>TakeIt – Full-Stack Developer & Platform Architect</b> | React, Node.js, Express, PostgreSQL, REST APIs
    </div>
    <ul class="bullets">
      <li>Architected an end-to-end event management and ticketing platform delivering fluid seat reservations, organizer administration consoles, and automated QR pass verification.</li>
      <li>Engineered transactional database locking mechanisms in PostgreSQL to eliminate duplicate ticket allocations and race conditions under concurrent booking traffic.</li>
      <li>Built responsive client interfaces with real-time seat inventory indicators, event scheduling workflows, and secure payment gateway integrations.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>Bomberman BattleRoyale – Client-Server Systems Architect & Game Developer</b> | Java, Multi-Threading, TCP/IP Sockets, Swing, OOP
    </div>
    <ul class="bullets">
      <li>Engineered a 2D multiplayer client-server architecture utilizing custom TCP/IP socket protocols and binary serialization to synchronize multi-player state with minimal packet latency.</li>
      <li>Implemented strict Object-Oriented Programming (OOP) design patterns and concurrent multithreading models to decouple the continuous game loop and rendering engine from network I/O.</li>
      <li>Designed procedural arena destruction algorithms, deterministic grid collision detection, and authoritative server-side hit registration.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>Fusion Rush – Lead Developer & AI Systems Architect</b> | Python, JavaScript, Discrete Math, Gamification Engine
    </div>
    <ul class="bullets">
      <li>Architected an educational AI tutoring and gamified logic application verifying Discrete Mathematics Rules of Inference in real-time.</li>
      <li>Built a rule-based AST inference validation engine in Python, delivering immediate pedagogical feedback and reducing student error rates on symbolic logic exercises.</li>
    </ul>
  </div>

  <!-- Technical Skills -->
  <div class="section-title">TECHNICAL SKILLS</div>

  <div class="skills-group">
    <b>Frontend Development:</b> TypeScript, JavaScript (ES6+), React.js, Next.js, HTML5, CSS3, Tailwind CSS, Responsive Web Design, UI/UX Design (Figma), PWA
  </div>
  <div class="skills-group">
    <b>Backend & API Engineering:</b> Node.js, Express.js, Python (FastAPI, Django), Java, C#, C++, C, PHP, RESTful APIs, WebSockets, TCP/IP Sockets, Multi-threading
  </div>
  <div class="skills-group">
    <b>Databases & State Management:</b> PostgreSQL, MongoDB, MySQL, Supabase, SQLite, Redis, Relational Modeling, Query Optimization, Database Locking, RLS
  </div>
  <div class="skills-group">
    <b>Cloud, DevOps & SRE:</b> AWS (EC2, S3, VPC, IAM, RDS), Docker & Containers, Kubernetes, Linux Server Administration, CI/CD Pipelines, Git, GitHub, Bash Scripting
  </div>
  <div class="skills-group">
    <b>Mobile & Payment Platforms:</b> React Native, Android Studio (Java/Kotlin), Stripe API Integration, Mobile-First Architecture, Memory Management
  </div>
  <div class="skills-group">
    <b>AI Engineering & Applied Tools:</b> AI-Assisted Tooling, Computer Vision, OCR Pipelines, Rule-Based Inference Engines, AST Parsing, Discrete Logic Verification
  </div>
  <div class="skills-group">
    <b>Software QA & Testing:</b> Test Planning, Regression Testing, Smoke Testing, Defect Tracking & Bug Fixing (Jira), API Testing (Postman), TDD
  </div>
  <div class="skills-group">
    <b>Foundational Computer Science:</b> Data Structures and Algorithms, Automata Theory & Formal Languages, Systems Architecture, Distributed Systems
  </div>
  <div class="skills-group">
    <b>Spoken Languages:</b> Cebuano (Native), English (Fluent), Filipino (Fluent)
  </div>

  <!-- Education -->
  <div class="section-title">EDUCATION</div>

  <div class="entry">
    <div class="row-between">
      <div class="title-left">Cebu Institute of Technology - University – BS Computer Science</div>
      <div class="date-right">2022 – 2026</div>
    </div>
  </div>
  <div class="entry" style="margin-top: -3pt;">
    <div class="row-between">
      <div class="title-left">Philippine Science High School - Central Visayas Campus – STEM Strand (With Honors)</div>
      <div class="date-right">2016 – 2022</div>
    </div>
  </div>

  <!-- Research Conferences & Competitions -->
  <div class="section-title" style="page-break-before: always;">RESEARCH CONFERENCES & COMPETITIONS</div>

  <div class="entry">
    <div class="entry-header">
      <b>Primary Author & Software Architect, International Conference on Education and Training Technologies (ICETT 2026)</b> | Scopus-Indexed
    </div>
    <div class="entry-desc">
      Architected and authored <i>"Fusion Rush: Enhancing Proficiency in Boolean Logic Algebra through Gamification"</i>. Engineered an adaptive AI tutoring software and gamified inference verification system, conducting empirical classroom deployments across Discrete Mathematics cohorts at CIT-U College of Computer Studies.
    </div>
  </div>

  <div class="entry">
    <div class="entry-header">
      <b>Author & Presenter, 10th Philippine Computing Science Congress (PCSC 2026)</b> | Davao
    </div>
    <div class="entry-desc">
      Developed and presented <i>"Bridging the Gap Between Synthetic and Real Handwriting: Using Visual AI to Accurately Parse Structured Logic Proofs"</i>. Engineered an end-to-end computer vision and structural parsing pipeline converting handwritten Discrete Math logic proofs into structured Abstract Syntax Trees (ASTs) for automated grading and semantic validity verification.
    </div>
  </div>

  <!-- Certifications -->
  <div class="section-title">CERTIFICATIONS</div>

  <div class="cert-row">
    <span class="cert-name">PhilNITS Fundamental Information Technology Engineers (FE) Certification</span>
    <span class="cert-date">Apr. 2025</span>
  </div>
  <div class="cert-row">
    <span class="cert-name">Level 3 TOPCIT Passer (Test of Practical Competency in IT)</span>
    <span class="cert-date">Jun. 2025</span>
  </div>
  <div class="cert-row">
    <span class="cert-name">AWS Academy Graduate - Cloud Architecting - Training Badge</span>
    <span class="cert-date">Nov. 2025</span>
  </div>
  <div class="cert-row">
    <span class="cert-name">AWS Academy Graduate - Cloud Foundations - Training Badge</span>
    <span class="cert-date">Nov. 2025</span>
  </div>
  <div class="cert-row">
    <span class="cert-name">CodeChum Java Certification</span>
    <span class="cert-date">Mar. 2024</span>
  </div>
  <div class="cert-row">
    <span class="cert-name">CodeChum C Certification</span>
    <span class="cert-date">Apr. 2023</span>
  </div>

  <!-- Volunteering and Leadership Experience -->
  <div class="section-title">VOLUNTEERING AND LEADERSHIP EXPERIENCE</div>

  <div class="entry">
    <div class="row-between">
      <div class="title-left">Koro Teknoy Chorale – Presidential OIC & Bass 2</div>
      <div class="date-right">Jun. 2022 – Present</div>
    </div>
    <ul class="bullets">
      <li>Served as Presidential Officer-in-Charge directing choral operations, institutional representation, and rehearsal governance for the university’s premier choral ensemble.</li>
      <li>Coordinated stage logistics, member scheduling, and vocal performances for major academic ceremonies and institutional university assemblies.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="row-between">
      <div class="title-left">Concert for a Cause – Featured Choral Performer</div>
      <div class="date-right">Jan. 2023</div>
    </div>
    <ul class="bullets">
      <li>Volunteered as featured vocalist performing in the Bradford Grand Choir benefit concert, mobilizing fundraising and relief donations for families displaced by Typhoon Odette.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="row-between">
      <div class="title-left">NSTP-CWTS Labangon Outreach – Community Welfare Volunteer</div>
      <div class="date-right">Apr. 2023</div>
    </div>
    <ul class="bullets">
      <li>Organized direct logistics, supply sorting, and donation distribution campaigns serving persons with disabilities (PWDs) across the Labangon community.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="row-between">
      <div class="title-left">Household of Faith Choir – Featured Sacred Music Vocalist</div>
      <div class="date-right">2022</div>
    </div>
    <ul class="bullets">
      <li>Performed as choral vocalist for major seasonal productions including the annual Easter Cantata and Christmas Cantata.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="row-between">
      <div class="title-left">Red Hot Pawns Chess Club – Active Member & Scholastic Player</div>
      <div class="date-right">2018 – 2020</div>
    </div>
    <ul class="bullets">
      <li>Represented the campus in inter-scholastic chess tournaments and co-facilitated chess strategy coaching sessions at Philippine Science High School - CVC.</li>
    </ul>
  </div>

</body>
</html>
"""
    return html_content

def compile_pdf_via_chrome(html_path, output_pdf_path):
    chrome_paths = [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    ]
    browser = None
    for p in chrome_paths:
        if os.path.exists(p):
            browser = p
            break
            
    if not browser:
        raise RuntimeError("Neither Chrome nor Edge was found for PDF rendering.")
        
    cmd = [
        browser,
        "--headless=new",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--print-to-pdf={output_pdf_path}",
        html_path
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"Error compiling PDF via browser: {res.stderr}")
        raise RuntimeError(res.stderr)
    print(f"Successfully generated PDF via {os.path.basename(browser)}: {output_pdf_path}")

def add_bottom_border_to_paragraph(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '000000')
    pbdr.append(bottom)
    pPr.append(pbdr)

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = parse_xml(f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    rPr = parse_xml(f'<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:rFonts w:ascii="Garamond" w:hAnsi="Garamond"/><w:color w:val="1155CC"/><w:u w:val="single"/></w:rPr>')
    new_run.append(rPr)
    text_elem = parse_xml(f'<w:t xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">{text}</w:t>')
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def create_docx_cv(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.38)
        section.bottom_margin = Inches(0.38)
        section.left_margin = Inches(0.42)
        section.right_margin = Inches(0.42)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Garamond'
    normal_style.font.size = Pt(9.8)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.space_before = Pt(0)

    # Header
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(1)
    run_name = name_p.add_run("Roddneil B. Gemina")
    run_name.font.name = "Garamond"
    run_name.font.size = Pt(20)
    run_name.font.bold = True

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(2)
    run_c = contact_p.add_run("roddneilgemina@gmail.com | +63 956-241-3577 | Cebu City, PH")
    run_c.font.name = "Garamond"
    run_c.font.size = Pt(10)

    links_p = doc.add_paragraph()
    links_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links_p.paragraph_format.space_after = Pt(5)
    run_l = links_p.add_run("github.com/RoddneilGemina | linkedin.com/in/roddneil-gemina-20b423409")
    run_l.font.name = "Garamond"
    run_l.font.size = Pt(10)

    def add_sec_title(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = "Garamond"
        r.font.size = Pt(11)
        r.font.bold = True
        add_bottom_border_to_paragraph(p)
        return p

    def add_bullet(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(1.8)
        p.paragraph_format.line_spacing = 1.15
        
        sq = p.add_run("■  ")
        sq.font.name = "Garamond"
        sq.font.size = Pt(6.8)
        sq.font.bold = True
        
        run = p.add_run(text)
        run.font.name = "Garamond"
        run.font.size = Pt(9.6)
        return p

    def add_bullet_with_bold_prefix(prefix_bold, text_rest):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(1.8)
        p.paragraph_format.line_spacing = 1.15
        
        sq = p.add_run("■  ")
        sq.font.name = "Garamond"
        sq.font.size = Pt(6.8)
        sq.font.bold = True
        
        r_pre = p.add_run(prefix_bold)
        r_pre.font.name = "Garamond"
        r_pre.font.size = Pt(9.6)
        r_pre.font.bold = True
        
        run = p.add_run(text_rest)
        run.font.name = "Garamond"
        run.font.size = Pt(9.6)
        return p

    def add_two_column_line(left_bold, left_rest, right_text, space_before=Pt(3), space_after=Pt(1)):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = space_before
        p.paragraph_format.space_after = space_after
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.43), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.keep_with_next = True
        
        r1 = p.add_run(left_bold)
        r1.font.name = "Garamond"
        r1.font.size = Pt(10.2)
        r1.font.bold = True
        
        if left_rest:
            r2 = p.add_run(left_rest)
            r2.font.name = "Garamond"
            r2.font.size = Pt(10.2)
            
        r3 = p.add_run(f"\t{right_text}")
        r3.font.name = "Garamond"
        r3.font.size = Pt(9.8)
        return p

    # 1. PROFESSIONAL EXPERIENCE
    add_sec_title("PROFESSIONAL EXPERIENCE")

    def add_exp_entry(role_company, location_rest, date_str, bullets):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(role_company)
        r.font.bold = True
        r.font.size = Pt(10.2)
        if location_rest:
            r2 = p.add_run(f" | {location_rest}")
            r2.font.size = Pt(10.2)

        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_after = Pt(1.5)
        p_d.paragraph_format.keep_with_next = True
        r_d = p_d.add_run(date_str)
        r_d.font.italic = True
        r_d.font.size = Pt(9.8)

        for b in bullets:
            if isinstance(b, tuple):
                add_bullet_with_bold_prefix(b[0], b[1])
            else:
                add_bullet(b)

    add_exp_entry(
        "Opswerks – Site Reliability Engineering (SRE) Intern",
        "Cebu City, PH",
        "January 2026 – Present",
        [
            "Administered enterprise Linux server environments, configuring AWS cloud infrastructure, containerized microservices via Docker, and production-grade Kubernetes pod orchestration.",
            "Constructed automated deployment pipelines and integrated cluster telemetry, distributed logging, and health monitoring tools to sustain high platform reliability and zero-downtime operations.",
            "Automated repetitive operational runbooks and system provisioning using Python and Bash scripting, significantly reducing manual incident mitigation time and operational toil."
        ]
    )

    add_exp_entry(
        "DXForm Technologies – Quality Assurance Analyst & Software Engineer",
        "Cebu City, PH",
        "January 2025 – August 2025",
        [
            ("Gigi POS (Next.js, React Native, PostgreSQL, Stripe): ", "Directed core QA testing strategies, authoring structured test matrices and regression suites across web and mobile POS terminals; contributed directly to UI/UX design, application logic, and product experience while developing code-level bug fixes and ensuring transactional integrity for Stripe payment processing."),
            ("Aicquire (React, Node.js, MongoDB, AI-Assisted Tooling): ", "Spearheaded frontend UI/UX architecture and responsive components in React, engineered backend application logic and MongoDB data pipelines in Node.js, and refined product experience to streamline AI-assisted tooling interactions."),
            "Logged reproducible defect lifecycles with root-cause analysis in Jira, collaborated cross-functionally with dev leads to optimize REST API response payloads, and enforced cross-platform UX fidelity."
        ]
    )

    add_exp_entry(
        "Nexchef – Mobile Architect & Full-Stack Developer (Contract / Freelance)",
        "Cebu City, PH",
        "November 2024 – March 2025",
        [
            "Developed a full-stack culinary mobile application in React Native and TypeScript featuring modular recipe creation, live step-by-step guidance, and synchronized multi-stage culinary timer engines.",
            "Integrated Supabase Realtime and WebSockets for instant state synchronization across multiple client devices, backed by PostgreSQL relational persistence.",
            "Designed responsive multi-screen layouts with micro-interactions, local offline persistence, and low-latency client-side caching."
        ]
    )

    add_exp_entry(
        "Refertoire – Systems Architect & Full-Stack Developer (Contract / Freelance)",
        "Cebu City, PH",
        "January 2026 – Present",
        [
            "Architected a 100% offline-capable music repertoire management platform, allowing performing ensembles to organize, search, and view complete digital sheet music libraries without internet dependency.",
            "Engineered local offline persistence utilizing IndexedDB and modern Service Workers, coupled with bidirectional client synchronization protocols upon network reconnection.",
            "Implemented responsive high-resolution document viewing and instant multi-parameter filtering across hundreds of musical scores."
        ]
    )

    add_exp_entry(
        "DishCover – Full-Stack Developer & Database Architect (Contract / Freelance)",
        "Cebu City, PH",
        "May 2024 – October 2024",
        [
            "Developed a comprehensive pantry and household inventory management application featuring automated expiration tracking, replenishment alerts, and storage categorization.",
            "Engineered normalized PostgreSQL relational schemas optimized for high-volume item lookup, barcode data association, and historical consumption analytics.",
            "Built a responsive, minimalist UI with Tailwind CSS and Next.js server components, optimizing page loads and delivering sub-second interaction speed."
        ]
    )

    add_exp_entry(
        "CropConnect – Full-Stack Agritech Developer (Contract / Freelance)",
        "Cebu City, PH",
        "February 2024 – August 2024",
        [
            "Co-developed an agricultural direct-to-consumer e-commerce platform connecting regional farmers directly with commercial buyers to eliminate intermediary price gouging.",
            "Engineered catalog browsing, inventory batching, and order fulfillment tracking pipelines with PostgreSQL backend models and responsive UI flows."
        ]
    )

    doc.add_page_break()
    # 2. RELEVANT PROJECTS
    add_sec_title("RELEVANT PROJECTS")

    def add_project(name_role, tech_str, bullets):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3.5)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(name_role)
        r.font.bold = True
        r.font.size = Pt(10.2)
        r2 = p.add_run(f" | {tech_str}")
        r2.font.size = Pt(10.2)
        for b in bullets:
            add_bullet(b)

    add_project(
        "ResuTex – Lead Architect & Full-Stack Developer",
        "Next.js 16, Python, ReportLab, Gemini Flash, Supabase, REST APIs",
        [
            "Architected an autonomous ATS resume studio combining Python ReportLab and LaTeX engines to dynamically compile and format tailored resumes and cover letters.",
            "Engineered dynamic ATS match scoring and qualification extraction algorithms against job postings, integrating LLM-assisted tailoring pipelines with Google Gemini Flash.",
            "Integrated Supabase REST APIs and local failover ledgers for real-time application tracking, status telemetry, and automated cold outreach drafting."
        ]
    )

    add_project(
        "TakeIt – Full-Stack Developer & Platform Architect",
        "React, Node.js, Express, PostgreSQL, REST APIs",
        [
            "Architected an end-to-end event management and ticketing platform delivering fluid seat reservations, organizer administration consoles, and automated QR pass verification.",
            "Engineered transactional database locking mechanisms in PostgreSQL to eliminate duplicate ticket allocations and race conditions under concurrent booking traffic.",
            "Built responsive client interfaces with real-time seat inventory indicators, event scheduling workflows, and secure payment gateway integrations."
        ]
    )

    add_project(
        "Bomberman BattleRoyale – Client-Server Systems Architect & Game Developer",
        "Java, Multi-Threading, TCP/IP Sockets, Swing, OOP",
        [
            "Engineered a 2D multiplayer client-server architecture utilizing custom TCP/IP socket protocols and binary serialization to synchronize multi-player state with minimal packet latency.",
            "Implemented strict Object-Oriented Programming (OOP) design patterns and concurrent multithreading models to decouple the continuous game loop and rendering engine from network I/O.",
            "Designed procedural arena destruction algorithms, deterministic grid collision detection, and authoritative server-side hit registration."
        ]
    )

    add_project(
        "Fusion Rush – Lead Developer & AI Systems Architect",
        "Python, JavaScript, Discrete Math, Gamification Engine",
        [
            "Architected an educational AI tutoring and gamified logic application verifying Discrete Mathematics Rules of Inference in real-time.",
            "Built a rule-based AST inference validation engine in Python, delivering immediate pedagogical feedback and reducing student error rates on symbolic logic exercises."
        ]
    )

    # 3. TECHNICAL SKILLS
    add_sec_title("TECHNICAL SKILLS")

    skills_data = [
        ("Frontend Development", "TypeScript, JavaScript (ES6+), React.js, Next.js, HTML5, CSS3, Tailwind CSS, Responsive Web Design, UI/UX Design (Figma), PWA"),
        ("Backend & API Engineering", "Node.js, Express.js, Python (FastAPI, Django), Java, C#, C++, C, PHP, RESTful APIs, WebSockets, TCP/IP Sockets, Multi-threading"),
        ("Databases & State Management", "PostgreSQL, MongoDB, MySQL, Supabase, SQLite, Redis, Relational Modeling, Query Optimization, Database Locking, RLS"),
        ("Cloud, DevOps & SRE", "AWS (EC2, S3, VPC, IAM, RDS), Docker & Containers, Kubernetes, Linux Server Administration, CI/CD Pipelines, Git, GitHub, Bash Scripting"),
        ("Mobile & Payment Platforms", "React Native, Android Studio (Java/Kotlin), Stripe API Integration, Mobile-First Architecture, Memory Management"),
        ("AI Engineering & Applied Tools", "AI-Assisted Tooling, Computer Vision, OCR Pipelines, Rule-Based Inference Engines, AST Parsing, Discrete Logic Verification"),
        ("Software QA & Testing", "Test Planning, Regression Testing, Smoke Testing, Defect Tracking & Bug Fixing (Jira), API Testing (Postman), TDD"),
        ("Foundational Computer Science", "Data Structures and Algorithms, Automata Theory & Formal Languages, Systems Architecture, Distributed Systems"),
        ("Spoken Languages", "Cebuano (Native), English (Fluent), Filipino (Fluent)")
    ]

    for cat, items in skills_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2.2)
        p.paragraph_format.line_spacing = 1.15
        r_b = p.add_run(f"{cat}: ")
        r_b.font.bold = True
        r_b.font.size = Pt(9.5)
        r_t = p.add_run(items)
        r_t.font.size = Pt(9.5)

    # 4. EDUCATION
    add_sec_title("EDUCATION")
    add_two_column_line("Cebu Institute of Technology - University", " – BS Computer Science", "2022 – 2026")
    add_two_column_line("Philippine Science High School - Central Visayas Campus", " – STEM Strand (With Honors)", "2016 – 2022", space_before=Pt(2))

    doc.add_page_break()
    # 5. RESEARCH CONFERENCES & COMPETITIONS
    add_sec_title("RESEARCH CONFERENCES & COMPETITIONS")

    def add_research_entry(title_bold, title_rest, desc_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3.5)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_bold)
        r.font.bold = True
        r.font.size = Pt(10.2)
        if title_rest:
            r2 = p.add_run(f" | {title_rest}")
            r2.font.size = Pt(10.2)
            
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(3.5)
        p_desc.paragraph_format.line_spacing = 1.15
        r_desc = p_desc.add_run(desc_text)
        r_desc.font.size = Pt(9.6)

    add_research_entry(
        "Primary Author & Software Architect, International Conference on Education and Training Technologies (ICETT 2026)",
        "Scopus-Indexed",
        "Architected and authored 'Fusion Rush: Enhancing Proficiency in Boolean Logic Algebra through Gamification'. Engineered an adaptive AI tutoring software and gamified inference verification system, conducting empirical classroom deployments across Discrete Mathematics cohorts at CIT-U College of Computer Studies."
    )

    add_research_entry(
        "Author & Presenter, 10th Philippine Computing Science Congress (PCSC 2026)",
        "Davao",
        "Developed and presented 'Bridging the Gap Between Synthetic and Real Handwriting: Using Visual AI to Accurately Parse Structured Logic Proofs'. Engineered an end-to-end computer vision and structural parsing pipeline converting handwritten Discrete Math logic proofs into structured Abstract Syntax Trees (ASTs) for automated grading and semantic validity verification."
    )

    # 6. CERTIFICATIONS
    add_sec_title("CERTIFICATIONS")

    certs = [
        ("PhilNITS Fundamental Information Technology Engineers (FE) Certification", "Apr. 2025"),
        ("Level 3 TOPCIT Passer (Test of Practical Competency in IT)", "Jun. 2025"),
        ("AWS Academy Graduate - Cloud Architecting - Training Badge", "Nov. 2025"),
        ("AWS Academy Graduate - Cloud Foundations - Training Badge", "Nov. 2025"),
        ("CodeChum Java Certification", "Mar. 2024"),
        ("CodeChum C Certification", "Apr. 2023")
    ]

    for c_title, c_date in certs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.8)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.43), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
        r1 = p.add_run(c_title)
        r1.font.bold = True
        r1.font.size = Pt(9.8)
        r2 = p.add_run(f"\t{c_date}")
        r2.font.size = Pt(9.8)

    # 7. VOLUNTEERING AND LEADERSHIP EXPERIENCE
    add_sec_title("VOLUNTEERING AND LEADERSHIP EXPERIENCE")

    def add_volunteer_entry(role_bold, date_str, bullets):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3.5)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.43), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.keep_with_next = True
        r1 = p.add_run(role_bold)
        r1.font.bold = True
        r1.font.size = Pt(10.2)
        r2 = p.add_run(f"\t{date_str}")
        r2.font.size = Pt(9.8)
        for b in bullets:
            add_bullet(b)

    add_volunteer_entry(
        "Koro Teknoy Chorale – Presidential OIC & Bass 2",
        "Jun. 2022 – Present",
        [
            "Served as Presidential Officer-in-Charge directing choral operations, institutional representation, and rehearsal governance for the university’s premier choral ensemble.",
            "Coordinated stage logistics, member scheduling, and vocal performances for major academic ceremonies and institutional university assemblies."
        ]
    )

    add_volunteer_entry(
        "Concert for a Cause – Featured Choral Performer",
        "Jan. 2023",
        [
            "Volunteered as featured vocalist performing in the Bradford Grand Choir benefit concert, mobilizing fundraising and relief donations for families displaced by Typhoon Odette."
        ]
    )

    add_volunteer_entry(
        "NSTP-CWTS Labangon Outreach – Community Welfare Volunteer",
        "Apr. 2023",
        [
            "Organized direct logistics, supply sorting, and donation distribution campaigns serving persons with disabilities (PWDs) across the Labangon community."
        ]
    )

    add_volunteer_entry(
        "Household of Faith Choir – Featured Sacred Music Vocalist",
        "2022",
        [
            "Performed as choral vocalist for major seasonal productions including the annual Easter Cantata and Christmas Cantata."
        ]
    )

    add_volunteer_entry(
        "Red Hot Pawns Chess Club – Active Member & Scholastic Player",
        "2018 – 2020",
        [
            "Represented the campus in inter-scholastic chess tournaments and co-facilitated chess strategy coaching sessions at Philippine Science High School - CVC."
        ]
    )

    doc.save(output_path)
    print(f"Successfully generated DOCX: {output_path}")

if __name__ == '__main__':
    html_str = create_html_cv()
    html_file = os.path.abspath("Gemina_CV_Full-Stack.html")
    with open(html_file, "w", encoding="utf-8") as f:
        f.write(html_str)
    print(f"Wrote HTML to {html_file}")
    
    pdf_file = os.path.abspath("Gemina_CV_Full-Stack.pdf")
    compile_pdf_via_chrome(html_file, pdf_file)
    
    docx_file = os.path.abspath("Gemina_CV_Full-Stack.docx")
    create_docx_cv(docx_file)
    
    dl_pdf = os.path.expanduser("~/Downloads/Gemina_CV_Full-Stack.pdf")
    dl_docx = os.path.expanduser("~/Downloads/Gemina_CV_Full-Stack.docx")
    
    shutil.copyfile(pdf_file, dl_pdf)
    shutil.copyfile(docx_file, dl_docx)
    print(f"Copied to Downloads: {dl_pdf} and {dl_docx}")
